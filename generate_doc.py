from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

doc = Document()

# --- Styles ---
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)
style.paragraph_format.space_after = Pt(6)

def add_heading(text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    return h

def add_bold_para(bold_text, normal_text=""):
    p = doc.add_paragraph()
    run = p.add_run(bold_text)
    run.bold = True
    if normal_text:
        p.add_run(normal_text)
    return p

def add_bullet(text, bold_prefix=""):
    p = doc.add_paragraph(style="List Bullet")
    if bold_prefix:
        run = p.add_run(bold_prefix)
        run.bold = True
        p.add_run(text)
    else:
        p.add_run(text)
    return p

def add_qa(question, answer):
    p = doc.add_paragraph()
    q_run = p.add_run(question)
    q_run.bold = True
    q_run.font.color.rgb = RGBColor(0x1E, 0x40, 0xAF)
    p.add_run("\n" + answer)
    p.paragraph_format.space_after = Pt(10)

# ============================================================
# TITLE
# ============================================================
title = doc.add_heading("TrustReply Demo — Discussion Guide", level=0)
for run in title.runs:
    run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

p = doc.add_paragraph()
run = p.add_run("Call with Josh McKibben, Head of Trust @ Harvey AI")
run.font.size = Pt(13)
run.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

p2 = doc.add_paragraph()
run2 = p2.add_run("Demo: https://trustreply-demo.vercel.app/")
run2.font.color.rgb = RGBColor(0x3B, 0x82, 0xF6)
run2.underline = True

doc.add_paragraph("")

# ============================================================
# OPENER
# ============================================================
add_heading("Opening (2 min)", level=1)
p = doc.add_paragraph()
p.add_run('"Thanks for making time, Josh. I built this as a working app — not slides — because I wanted to show how I think through a problem end to end: data model first, then product layer, then agent architecture. Let me share my screen and walk you through it live."').italic = True

# ============================================================
# 5 KEY TALKING POINTS
# ============================================================
add_heading("The 5 Things You Must Land", level=1)

add_bold_para("1. You started with the data model, not the UI")
doc.add_paragraph(
    '"I worked backwards from your brief — compliance data lake, tables, schemas, views. '
    'The 8-table Snowflake schema is the foundation. Everything else — the dashboard, the agents, '
    'the pipeline — sits on top of that data model. I think the data model is the product decision that matters most."'
)

add_bold_para("2. The flywheel is the real product insight")
doc.add_paragraph(
    '"Switch to Goldman Sachs in the dropdown. Now Allen & Overy. Watch the reuse rate: '
    '84% → 93% → 96%. Every approved answer makes the next questionnaire faster. '
    "That's not a feature — that's the business case. The system gets smarter with each customer.\""
)

add_bold_para("3. Test once, audit many")
doc.add_paragraph(
    '"When the pipeline hits Stage 4, CC-05 lights up in the control grid. One answer to one encryption question '
    "satisfies SOC 2, ISO 27001, AND PCI DSS. That's your leverage — 16 control families that map across every "
    'framework Harvey holds. The framework badges on each card make that visible."'
)

add_bold_para("4. Trust scoring is the core product question")
doc.add_paragraph(
    '"Every AI output carries a Trust Index — confidence, accuracy, freshness, benchmark — '
    "banded as TRUSTED, REVIEW, or FLAG. This is the answer to 'Should I trust this AI-generated response?' "
    'High confidence auto-routes to one-click approval. Low confidence escalates to a senior analyst. '
    'The human stays in the loop where it matters."'
)

add_bold_para("5. Buy vs Build maps to his architecture vision")
doc.add_paragraph(
    '"In the Data Architecture tab — Anecdotes is the buy, everything to the right is the build. '
    "Snowpipe for ingestion, Cortex for inference inside Snowflake so no data leaves the warehouse. "
    'This maps directly to what you described in your email."'
)

# ============================================================
# DEMO FLOW TABLE
# ============================================================
add_heading("Demo Flow (15-20 min)", level=1)

table = doc.add_table(rows=10, cols=3)
table.style = "Light Grid Accent 1"
table.alignment = WD_TABLE_ALIGNMENT.CENTER

headers = ["What to Show", "What to Say", "Time"]
for i, h in enumerate(headers):
    cell = table.rows[0].cells[i]
    cell.text = h
    for p in cell.paragraphs:
        for r in p.runs:
            r.bold = True

rows_data = [
    ("Dashboard — let pipeline auto-play", '"This traces one real questionnaire through 8 stages"', "2 min"),
    ("Pause at Stage 4, point to CC-05 highlight", '"Test once, audit many — one answer, three frameworks"', "1 min"),
    ("Client dropdown → Goldman → Allen & Overy", '"The flywheel — reuse climbs from 84% to 96%"', "2 min"),
    ("Hit Compare button", '"Side by side, the compounding is obvious"', "1 min"),
    ("Mission Control — 3-tier agent architecture", '"Specialized agents, not monolithic. Quality tier keeps it honest."', "2 min"),
    ("TrustReply Agent chat", '"Every response scored. TRUSTED, REVIEW, or FLAG."', "1 min"),
    ("Data Model — interactive ERD", '"This is the compliance data lake — 8 tables, proper relationships"', "2 min"),
    ("Data Architecture — 5-layer flow", '"Buy the GRC layer, build the intelligence layer. Your words."', "3 min"),
    ("How I Built This (skim)", '"My design decisions and open questions — things I know I got wrong"', "1 min"),
]

for i, (col1, col2, col3) in enumerate(rows_data):
    row = table.rows[i + 1]
    row.cells[0].text = col1
    row.cells[1].text = col2
    row.cells[2].text = col3

doc.add_paragraph("")

# ============================================================
# QUESTIONS TO ASK
# ============================================================
add_heading("Questions to Ask Josh", level=1)

add_bold_para("About the problem:")
add_bullet('"How many questionnaires is your team processing per quarter right now?"')
add_bullet('"What\'s the current turnaround — days? Weeks?"')
add_bullet('"What\'s the biggest pain point — volume, evidence staleness, audit prep, or something else?"')

add_bold_para("About the architecture:")
add_bullet('"Is Anecdotes the primary GRC tool?"')
add_bullet('"How mature is the Snowflake environment? Snowpipe running yet?"')
add_bullet('"Does the 6-agent breakdown map to how your team divides work?"')

add_bold_para("About the role:")
add_bullet('"What does the team look like today? Who would I work with?"')
add_bullet('"Where does this role sit — product management, data engineering, or hybrid?"')
add_bullet('"What does success look like in 6 months?"')

# ============================================================
# NEXT STEPS PITCH
# ============================================================
add_heading('Your "Next Steps" Pitch', level=1)

p = doc.add_paragraph()
p.add_run('"This demo is a prototype. Here\'s my 90-day plan to make it real:"').italic = True

doc.add_paragraph("")
add_bold_para("Month 1: Listen. ", "Shadow your analysts, map real workflows, learn what I got wrong.")
add_bold_para("Month 2: Connect real data. ", "Snowpipe from Anecdotes, evidence flowing through RAW → CURATED → ANALYTICS.")
add_bold_para("Month 3: Ship the first real agent. ", "RAG over actual evidence via Cortex, real confidence scoring.")

doc.add_paragraph(
    '"I\'m a PM who can prototype and speak data — I\'ve done this 0-to-1 pattern at EY and Intuit. '
    "I'd own product requirements and stakeholder alignment while partnering with your engineers on the infrastructure.\""
)

# ============================================================
# HANDLING TOUGH QUESTIONS
# ============================================================
add_heading("Handling Tough Questions", level=1)

add_qa(
    '"Why should we hire a PM who doesn\'t have security experience?"',
    '"My value isn\'t security domain expertise — it\'s that I can ship a data platform from zero. '
    "I did it at EY with ESG (14 metrics, calculation engine, full platform) and at Intuit migrating "
    "80+ analysts from Excel to a web forecasting tool. The pattern is the same: understand the domain, "
    "design the data model, build the product layer. I'd learn compliance on the job — and I learn fast, "
    'as this demo hopefully shows."'
)

add_qa(
    '"We could just buy a tool for this."',
    '"You should buy the GRC workflow layer — that\'s Anecdotes. What you can\'t buy is the intelligence layer '
    "that's specific to Harvey: your control framework, your evidence corpus, your trust scoring thresholds, "
    "your agent architecture tuned to your team's workflow. That's the build.\""
)

add_qa(
    '"What if the AI hallucinates?"',
    '"That\'s why the Trust Index exists. Every output is scored and routed — high confidence auto-approves, '
    "low confidence escalates. Plus the Quality tier agents catch drift and errors before they reach a customer. "
    'The system is designed to fail safe, not fail silent."'
)

add_qa(
    '"What would you change about this demo?"',
    '"Honestly — a lot. I built this from your brief and public information. I\'m certain my assumptions '
    "about your tooling, team structure, and priorities are partially wrong. That's the point of month 1 — "
    "I'd rather build something to react to than start with a blank page and a long discovery phase.\""
)

add_qa(
    '"This is impressive but it\'s just a frontend."',
    '"You\'re right — the data is mocked. But the architecture isn\'t hypothetical. The Snowflake schema is real, '
    "the Snowpipe/Streams/Tasks pattern is how I'd actually build the ingestion layer, and Cortex is a real "
    "Snowflake feature for in-warehouse LLM inference. The gap between this demo and production is data "
    'connectivity and engineering partnership — not a rethink."'
)

# ============================================================
# FLYWHEEL DATA
# ============================================================
add_heading("Flywheel Data — Client Comparison", level=1)

doc.add_paragraph("Use the client dropdown in the demo to show these numbers live:")

table2 = doc.add_table(rows=5, cols=4)
table2.style = "Light Grid Accent 1"
table2.alignment = WD_TABLE_ALIGNMENT.CENTER

headers2 = ["Metric", "Latham & Watkins (1st)", "Goldman Sachs (2nd)", "Allen & Overy (3rd)"]
for i, h in enumerate(headers2):
    cell = table2.rows[0].cells[i]
    cell.text = h
    for p in cell.paragraphs:
        for r in p.runs:
            r.bold = True

flywheel_data = [
    ("Duplicates found", "14 (5%)", "31 (17%)", "52 (35%)"),
    ("Reuse rate", "84.6%", "93.2%", "96.4%"),
    ("Review time", "12s", "8s", "6s"),
    ("Confidence", "0.96", "0.98", "0.97"),
]

for i, (col1, col2, col3, col4) in enumerate(flywheel_data):
    row = table2.rows[i + 1]
    row.cells[0].text = col1
    row.cells[1].text = col2
    row.cells[2].text = col3
    row.cells[3].text = col4

doc.add_paragraph("")

# ============================================================
# MINDSET REMINDERS
# ============================================================
add_heading("Mindset Reminders", level=1)

add_bullet("Listen more than you talk. ", "The demo speaks for itself. Your job is to ask questions and understand his real problems.")
add_bullet("Don't defend, learn. ", 'If he says something is wrong, say "Tell me more about that." His corrections are gold.')
add_bullet("Be honest about gaps. ", '"I don\'t know that yet, but here\'s how I\'d learn it" is stronger than faking expertise.')
add_bullet("Close with energy. ", '"I\'m genuinely excited about this problem. When can we talk next steps?"')

# ============================================================
# SAVE
# ============================================================
output_path = "/home/user/PMWorkOS/TrustReply_Josh_Call_Guide.docx"
doc.save(output_path)
print(f"Saved to {output_path}")
